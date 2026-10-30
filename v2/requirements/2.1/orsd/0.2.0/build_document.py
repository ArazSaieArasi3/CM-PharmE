import pathlib,json,copy
from docx import Document
from docx.shared import Pt,Inches,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
OUT=pathlib.Path(__file__).resolve().parent;OLD=OUT.parent/'0.1.0'
d=json.loads((OUT/'requirements.json').read_text());v=json.loads((OUT/'validation.json').read_text())
doc=Document(OLD/'CM_PharmE_2_1_ORSD_v0_1_0_FA.docx')
def paras():
 yield from doc.paragraphs
 for t in doc.tables:
  for row in t.rows:
   for c in row.cells:yield from c.paragraphs
 for s in doc.sections:yield from s.footer.paragraphs
def rtl(p):
 pr=p._p.get_or_add_pPr()
 for x in list(pr.findall(qn('w:bidi'))):pr.remove(x)
 e=OxmlElement('w:bidi');e.set(qn('w:val'),'1');pr.append(e)
 jc=pr.find(qn('w:jc'))
 if jc is None:jc=OxmlElement('w:jc');pr.append(jc)
 jc.set(qn('w:val'),'start')
 return p
def replace(p,text):
 p.text=text
 for r in p.runs:r.font.name='Vazirmatn';r.font.size=Pt(11)
 return p
def para(text):return rtl(doc.add_paragraph(text))
def head(text,level=1):return rtl(doc.add_heading(text,level))
for p in list(paras()):
 for run in p.runs:
  run.text=run.text.replace('0.1.0','0.2.0').replace('۲۰ اکتبر ۲۰۲۶','۲۱ اکتبر ۲۰۲۶').replace('۲۰ اکتبر','۲۱ اکتبر')
  run.text=run.text.replace('وضعیت پیشنهادی برای مرور علمی','مبنای تجمیع شده پژوهشی برای مرور محتوا')
 text=p.text
 if text.startswith('این سند نیازمندی‌های انتالوژی'):
  replace(p,'این سند مبنای تجمیع شده نیازمندی‌های CM-PharmE برای مسیر 2.1 است. نسخه 0.2.0 هم‌پوشانی‌های شناخته شده را تعیین تکلیف، نگاشت R-BA-02 را اصلاح و برای هر ردیف قرارداد پذیرش تعریف می‌کند. کنترل فنی سند و آزمون هدفمند انجام شده است؛ پذیرش علمی همه نیازها و اعتبارسنجی کامل انتالوژی از این وضعیت استنتاج نمی‌شود.')
 elif text.startswith('سیزده ردیف CL و CTX'):
  replace(p,'سیزده پالایش CL و CTX به نیازهای مرتبط متصل و از شمار نیازهای مستقل خارج نگه داشته شدند. ۲۴ تصمیم درباره هم‌پوشانی‌های شناخته شده ثبت شد. هشت بند به‌عنوان تخصص زمینه‌ای یک نیاز عمومی نگه داشته شدند؛ بندهای مکمل و موضوعات متمایز نیز ادغام نشدند. این تصمیم‌ها یک بازبینی تحریری و ردیابی برای مبنای پژوهش هستند و جایگزین داوری مستقل متخصص نیستند. فهرست تصمیم‌ها در بخش ۱۲ و پرونده overlap-decisions.json آمده است.')
 elif text.startswith('در اجرای تازه، ۱۲ کنترل'):
  replace(p,'در این نسخه ۲۲ کنترل سند و ردیابی، ۷ حالت پرسش تازه، ۳ حالت منطقی با HermiT، ۱۷ حالت پذیرش داده BA با SHACL و ۲۳ بازاجرای پرسش تاریخی مطابق انتظار بودند. این دسته‌ها با یکدیگر جمع نمی‌شوند و شمار نیازهای تأییدشده محسوب نمی‌شوند. دو جهش معیوب نیز عمداً ساخته شدند: نام رابطه قدیمی در پرسش و قید اجباری اجرای قابلیت در مدل؛ هر دو توسط آزمون مربوط شناسایی شدند.')
 elif text.startswith('یک اختلاف ردیابی مشخص شد'):
  replace(p,'اختلاف ردیابی R-BA-02 در دامنه فعلی برطرف شد: پرسش جدید به رابطه صادرشده capabilityRealizedInLogistics متصل است. نام عمومی capabilityRealizedIn هم‌ارز یا alias اعلام نشد و فایل‌های تاریخی حفظ شدند. پرسش روی حالت فاقد اجرای ثبت شده، دارای اجرا و گراف ترکیبی آزموده شد؛ مدل خلاف با سقف صفر تحقق سازگار و با جهش اجباری‌کردن تحقق ناسازگار بود. این نتیجه فقط مربوط به مسیر لجستیک در پروفایل بررسی شده است؛ نبود داده، نبود اجرای واقعی در جهان را ثابت نمی‌کند.')
 elif text.startswith('برای هر بند یک پرسش ارزیابی'):
  replace(p,'برای هر یک از ۸۳ ردیف، قرارداد پذیرش با ورودی، پاسخ موردانتظار، معیار مثبت و منفی، مرز داده نامعلوم و روش بررسی ثبت شد. ۷۱ قرارداد کارکردی و ۱۲ پروتکل کیفیتی داریم. ۲۴ بند چهار ماژول به قراردادهای محدود پیشین متصل‌اند؛ R-BA-02 آزمون تازه دارد. ۴۷ قرارداد کارکردی دیگر در این بسته اجرای پذیرش مستقل تازه ندارند. مشخص بودن قرارداد با گذراندن آن برابر نیست.')
 elif text.startswith('نسخه سند از نسخه انتالوژی مستقل است'):
  replace(p,'نسخه سند از نسخه انتالوژی مستقل است. ORSD 0.2.0 برای مسیر انتالوژی 2.1 یک مبنای تجمیع شده پژوهشی است و انتشار رسمی مدل ایجاد نمی‌کند. نسخه 0.1.0 بدون تغییر بایگانی شده است. پس از یک مرور محتوایی نویسنده، می‌توان همین مبنا را با ثبت اصلاحات لازم در 1.0.0 تثبیت کرد. تکمیل همه آزمون‌های مدل یا جمع‌آوری همه داده‌های تجربی شرط نوشتن و تثبیت مشخصات نیست؛ آن کارها در ارزیابی و تکامل انتالوژی ادامه می‌یابند.')
 elif text.startswith('مقایسه: نگاشت شناسه نیازمند حل'):
  replace(p,'مقایسه: ارجاع فعلی در بسته آزمایشی موجود است؛ اختلاف نام در سند حل شد. آزمون‌ها محدود به تحقق لجستیک هستند. پذیرش معنایی کامل مدل همچنان باز است.')

# Correct the individual requirement references in its own record.
inside=False
for p in doc.paragraphs:
 if p.style.name.startswith('Heading'):
  if p.text.startswith('R-BA-02'):inside=True
  elif inside:inside=False
 if inside and 'ارجاع مدل:' in p.text:replace(p,'Model reference: capabilityRealizedInLogistics')
 if inside and p.text.startswith('پرسش '):replace(p,'پرسش CQ-R-BA-02: کدام قابلیت در گراف بررسی شده فاقد اجرای لجستیک ثبت شده است؟')
 if inside and p.text.startswith('معیار مثبت:'):replace(p,'معیار مثبت: قابلیت بدون اجرای ثبت شده بازیابی شود و قابلیت دارای اجرای ثبت شده بازیابی نشود؛ جداگانه سازگاری مدل خلاف با صفر تحقق بررسی شود.')
 if inside and p.text.startswith('مقایسه:'):replace(p,'مقایسه: نگاشت جاری به capabilityRealizedInLogistics در بسته آزمایشی ثبت و آزموده شد؛ برای دامنه‌های دیگر تحقق، نگاشت جدا لازم است. پذیرش معنایی کامل انتالوژی باز است.')

for t in doc.tables:
 for row in list(t.rows):
  texts=[c.text for c in row.cells]
  if texts and texts[0]=='ارجاع موجود فقط در بسته آزمایشی':row.cells[1].text='22'
  if texts and texts[0]=='منابع فایل نسخه دار':row.cells[1].text=str(len(d['sources']))
  if texts and texts[0]=='نگاشت شناسه نیازمند حل':
   row.cells[0].text='نگاشت سندی حل نشده';row.cells[1].text='0';row.cells[2].text='اختلاف شناسایی شده R-BA-02 در دامنه لجستیک رفع شد؛ پذیرش معنا جداست.'
  if any('اولین تجمیع مستقل' in x for x in texts):row.cells[0].text='0.1.0';row.cells[-1].text='بایگانی شده؛ مبنای تاریخی بدون تغییر'
  if texts and 'برنامه ریزی شده' in texts[0]:
   row.cells[0].text='0.2.0';row.cells[1].text='تعیین تکلیف هم‌پوشانی‌ها، اصلاح نگاشت، ۸۳ قرارداد پذیرش و آزمون هدفمند';row.cells[2].text='نسخه حاضر؛ آماده مرور نهایی محتوا'
  if texts and texts[0]=='چهار ماژول مستقل':row.cells[2].text='نتایج اجرای تاریخی؛ بازاجراهای منتخب این نسخه جدا گزارش شده‌اند؛ پذیرش علمی کامل باز است.'
  if texts and texts[0]=='P0' and texts[1]=='مرز و هم‌پوشانی نیازها':
   row.cells[0].text='انجام شد';row.cells[2].text='۲۴ تصمیم ثبت شد؛ هشت تخصص زمینه‌ای و بندهای مکمل مشخص‌اند. مرور محتوایی نهایی باقی است.'
  if texts and texts[0]=='پذیرش ORSD 0.1':row.cells[0].text='تکمیل فنی ORSD 0.2';row.cells[1].text='ساختار، ردیابی، هم‌پوشانی‌ها و قراردادها کنترل شدند؛ یک مرور محتوایی برای تثبیت 1.0 لازم است.'
  for cell in row.cells:
   for p in cell.paragraphs:rtl(p)

def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for col,w in zip(t.columns,widths):col.width=Inches(w)
 props=t._tbl.tblPr;props.append(OxmlElement('w:bidiVisual'))
 borders=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  z=OxmlElement('w:'+side);z.set(qn('w:val'),'single');z.set(qn('w:sz'),'4');z.set(qn('w:color'),'D9D9D9');borders.append(z)
 props.append(borders)
 for c,h in zip(t.rows[0].cells,headers):c.text=h
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 for row in rows:
  for c,x in zip(t.add_row().cells,row):c.text=str(x)
 for i,row in enumerate(t.rows):
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for cell in row.cells:
   cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE6EF' if i==0 else 'F5F7F9' if i%2==0 else 'FFFFFF');cell._tc.get_or_add_tcPr().append(sh)
   for p in cell.paragraphs:
    rtl(p);p.paragraph_format.space_after=Pt(5)
    for r in p.runs:r.font.name='Vazirmatn';r.font.size=Pt(9.5);r.bold=i==0
 return t

head('۱۲ جمع بندی تصمیم ها و تثبیت سند')
para('در نسخه حاضر چهار کار فنی سند تکمیل شد: تعیین تکلیف هم‌پوشانی‌های شناخته شده، رفع اختلاف نگاشت، تعریف قراردادهای پذیرش و اجرای ممیزی و آزمون هدفمند. یک کار برای تثبیت 1.0 باقی است: مرور محتوایی نویسنده درباره هدف، مرز و اولویت‌ها و ثبت اصلاحات موردنظر. این مرور با ارزیابی مستقل متخصص دامنه یا اعتبارسنجی کل انتالوژی یکسان نیست.')
table(['کار سند','وضعیت','خروجی'],[['تعیین تکلیف هم پوشانی ها','انجام شد','۲۴ تصمیم؛ حفظ شناسه‌ها و هشت تخصص زمینه‌ای'],['رفع R BA 02','انجام شد','نگاشت جاری و پرسش با آزمون تمایز حالت اجرا'],['قراردادهای پذیرش','انجام شد','۷۱ کارکردی و ۱۲ کیفیتی؛ وضعیت اجرا برای هر بند روشن'],['کنترل فنی و خروجی','انجام شد','کنترل ساختار، داده آزمون و قالب Word و PDF'],['مرور محتوایی نویسنده','باقی مانده','مرور هدف و دامنه و اولویت؛ سپس تثبیت 1.0.0']],[2,1.2,3.8])
head('۱۲ ۱ تصمیم های هم پوشانی',2)
para('روابط عمومی و تخصصی زیر، سازمان‌دهی نیازمندی‌ها هستند. این برچسب‌ها به تنهایی ادعای entailment رسمی میان قیود OWL یا اثبات تعداد نیازهای معنایی مستقل ایجاد نمی‌کنند. هیچ ردیفی فقط برای بهتر شدن درصد پوشش حذف نشده است.')
table(['تصمیم','بندهای مرتبط','نتیجه و دلیل'],[[x['id'],'، '.join(x['requirement_ids']),x['rationale_fa']] for x in d['overlap_decisions']],[.7,1.8,4.5])
head('۱۲ ۲ نتیجه آزمون هدفمند',2)
table(['دسته','مطابق انتظار','معنا'],[['کنترل سند و ردیابی','22 / 22','کنترل شناسه، منبع، قرارداد و تصمیم'],['حالت های پرسش تازه','7 / 7','فاقد اجرا، دارای اجرا، ترکیبی، خالی، مرز نام قدیمی و کشف جهش پرسش'],['آزمون های منطقی','3 / 3','مدل خلاف صفر تحقق و حالت دارای اجرا سازگار؛ جهش اجرای اجباری ناسازگار'],['پذیرش داده BA','17 / 17','عدم پس‌رفت پذیرش نمونه‌های قبلی'],['پرسش های تاریخی','23 / 23','بازاجرای مرجع؛ شامل پرسش قدیمی که دیگر شاهد R-BA-02 نیست']],[2.2,1,3.8])
para('ردیف پرسش‌های تاریخی به معنی درست بودن همه پرسش‌های قدیمی در همه حالات نیست. نام قدیمی در گراف دارای اجرای جدید، پاسخ نادرست ایجاد کرد و همین شاهد حساسیت، دلیل جایگزینی قرارداد جاری است. ۴۷ قرارداد کارکردی دیگر هنوز در این بسته اجرای پذیرش مستقل تازه ندارند؛ این محدودیت در ارزیابی انتالوژی پیگیری می‌شود.')
head('۱۲ ۳ کار باقی مانده و مسیر نسخه بعد',2)
para('برای خود سند، افزایش بی‌دلیل تعداد نیازها یا تکرار این دور تدوین لازم نیست. مرور نهایی باید فقط به این موارد پاسخ دهد: آیا هدف و مرز درست است؛ آیا اولویت Must و Should برای تعهد پژوهش مناسب است؛ و آیا نیاز ضروری مشخصی جا افتاده است. پس از ثبت نتیجه این مرور، نسخه 1.0.0 را می‌توان به عنوان مبنای مشخصات تثبیت کرد و تغییرهای بعدی را با درخواست تغییر نسخه‌دار پیش برد.')
para('آزمون کامل نیازها، تکمیل داده‌های واقعی، مرور خبرگان و همگام‌سازی مقاله و دیاگرام در مسیر ارزیابی و انتشار انتالوژی ادامه دارند؛ لازم نیست تکمیل آن‌ها را با بسته شدن نسخه مشخصات یکی بدانیم. نسخه 0.2.0 فعلی آماده استفاده پژوهشی و مرور محتواست، با حفظ صریح حدود شواهد.')
para('مبنای سند پیشین در commit 32a964f26add298b6df25bf9896d81b69e83b9fe و مبنای مدل در commit 09d60136a990d6602438be4085a4d407c9bd5c68 ثبت است. چهار فایل ورودی افزوده، شامل requirements.json نسخه 0.1 و ontology.ttl و shapes.ttl و hierarchy.ttl ماژول BA، در sources.json با شناسه و hash آمده‌اند. هیچ فایل مدل یا شاهد تاریخی در این کار تغییر نکرده است.')
doc.core_properties.title='CM PharmE 2.1 Ontology Requirements Specification Document 0.2.0'
doc.core_properties.subject='Consolidated requirements baseline and targeted verification'
path=OUT/'CM_PharmE_2_1_ORSD_v0_2_0_FA.docx';doc.save(path)

md=['# CM PharmE 2.1 ORSD 0.2.0','','Consolidated research specification baseline | 21 October 2026','','## Scope of closure','','Technical preparation of the document is complete for the declared scope. One author content review remains before freezing 1.0.0. This is not independent expert acceptance or an ontology release.','','## Requirements','','83 retained tracked rows; 71 functional, 12 quality; eight explicit scoped specializations. Thirteen local refinements stay attached; historical B2-DS-07 remains merged into R-DS-04.','']
for x in d['requirements']:
 md += ['### '+x['id']+' '+x['domain'],'',x['statement_fa'],'',f"Disposition: {x['specification_disposition']}; parent: {x['specializes_requirement_id'] or 'none'}",f"Priority: {x['priority']}; sources: {', '.join(x['source_ids'])}",f"Question: {x.get('cq') or 'Quality protocol'}",f"Expected: {x['expected_answer']}",f"Positive: {x['positive_acceptance']}",f"Negative: {x['negative_acceptance']}",f"Anchors: {', '.join(x['anchors'])}",f"Model comparison: {x['comparison_status']}; semantic approval: {x['acceptance_status']}",'']
md+=['## Overlap decisions','']+[f"- {x['id']} {', '.join(x['requirement_ids'])}: {x['rationale_fa']}" for x in d['overlap_decisions']]+['','## Targeted verification','',json.dumps(v['summary'],ensure_ascii=False,indent=2),'','47 other functional contracts have no new standalone acceptance execution in this package. No semantic approval is inferred.','','## Next action','','One author review of purpose, scope, priorities and any concrete missing need, then freeze document 1.0.0. Broader ontology verification and expert evaluation continue separately.','']
(OUT/'ORSD.md').write_text('\n'.join(md))
print(path)
